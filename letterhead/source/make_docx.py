# -*- coding: utf-8 -*-
"""
Build the editable Word letterhead for Goya Ten.

The letterhead artwork lives in the *header* and *footer* of the A4 section, so
the page margins stay free for typing and the artwork repeats on every page.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Mm, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # .../letterhead
TMP = os.path.join(OUT, "source", "_build")

NAVY = RGBColor(0x0E, 0x2A, 0x47)
GRAY = RGBColor(0x6E, 0x76, 0x81)
DARK = RGBColor(0x33, 0x33, 0x33)


# ---------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_rtl_paragraph(p):
    """Right-to-left paragraph (bidi) + right alignment."""
    pPr = p._p.get_or_add_pPr()
    pPr.append(_el("w:bidi", **{"w:val": "1"}))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def make_rtl_run(run):
    rPr = run._r.get_or_add_rPr()
    rPr.append(_el("w:rtl", **{"w:val": "1"}))
    rPr.append(_el("w:cs", **{"w:val": "1"}))
    return run


def set_run_fonts(run, latin="Calibri", cs="Arial"):
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:cs"), cs)


def zero_para(p):
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return p


def para(doc, text="", size=11, color=DARK, bold=False, align=None, rtl=False,
         before=0, after=6, latin="Calibri", cs="Arial", italic=False):
    p = doc.add_paragraph()
    zero_para(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if rtl:
        set_rtl_paragraph(p)
    elif align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.italic = italic
        set_run_fonts(r, latin, cs)
        if rtl:
            make_rtl_run(r)
    return p


def add_picture_para(container, image, width_mm):
    """Place a full-bleed picture in a header/footer paragraph."""
    p = container.paragraphs[0]
    zero_para(p)
    p.paragraph_format.line_spacing = 1.0
    # let the artwork reach the paper edge (the text column is inset by the margins)
    p.paragraph_format.left_indent = Mm(-20)
    p.paragraph_format.right_indent = Mm(0)
    p.paragraph_format.first_line_indent = Mm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run()
    r.add_picture(image, width=Mm(width_mm))
    return p


# ---------------------------------------------------------------------------
def build_docx(path, header_png, footer_png):
    doc = Document()

    # ---- default style ----------------------------------------------------
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = DARK
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), "Calibri")
    rf.set(qn("w:hAnsi"), "Calibri")
    rf.set(qn("w:cs"), "Arial")
    rf.set(qn("w:eastAsia"), "Calibri")

    # ---- page setup -------------------------------------------------------
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(46)          # = header artwork height
    sec.bottom_margin = Mm(30.5)     # = footer artwork height
    sec.left_margin = Mm(20)
    sec.right_margin = Mm(20)
    sec.header_distance = Mm(0)
    sec.footer_distance = Mm(0)

    # ---- header / footer artwork -----------------------------------------
    art = Mm(210)                    # full bleed, exactly the page width
    add_picture_para(sec.header, header_png, 210)
    add_picture_para(sec.footer, footer_png, 210)

    # ---- ready-to-use letter body ----------------------------------------
    para(doc, "التاريخ / Date:  ................................",
         size=10.5, color=GRAY, rtl=True, after=14)

    para(doc, "السيد الأستاذ / .................................................",
         size=12, color=NAVY, bold=True, rtl=True, after=2)
    para(doc, "الموقر", size=12, color=NAVY, bold=True, rtl=True, after=10)

    para(doc, "تحية طيبة وبعد،،", size=11, rtl=True, after=10)

    for line in [
        "يسعدنا أن نتقدم إلى سيادتكم بخدمات شركة جويا تن، ونؤكد حرصنا على تقديم",
        "أفضل الحلول والمنتجات بأعلى معايير الجودة وفي التوقيتات المتفق عليها.",
    ]:
        para(doc, line, size=11, rtl=True, after=4)

    para(doc, "", after=6)
    para(doc, "Dear Sir / Madam,", size=11, after=8)
    para(doc, "Thank you for your interest in Goya Ten. Please find enclosed our latest "
              "offer; we remain at your disposal for any further information or "
              "clarification.", size=11, after=10)
    para(doc, "Yours faithfully,", size=11, after=26)

    para(doc, "المدير التنفيذي", size=11, color=NAVY, bold=True, rtl=True, after=2)
    para(doc, "جويا تن", size=12, color=NAVY, bold=True, rtl=True, after=14)
    para(doc, "Managing Director  -  Goya Ten", size=10.5, color=NAVY, bold=True)

    doc.save(path)
    return path


if __name__ == "__main__":
    h = f"{TMP}/word_header.png"
    f = f"{TMP}/word_footer.png"
    print(build_docx(f"{OUT}/Goya_Ten_Letterhead_Word.docx", h, f))
